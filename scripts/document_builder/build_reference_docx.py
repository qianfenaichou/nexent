#!/usr/bin/env python3
"""Build the Pandoc reference document used by ``build_docs.py``.

The template favors restrained typography, generous spacing, readable code
samples, and quiet navigation suitable for published product documentation.

Run: python scripts/build_reference_docx.py
Output: scripts/reference.docx
"""

from __future__ import annotations

from pathlib import Path


SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = SCRIPT_DIR / "reference.docx"

INK = "172033"
MUTED = "667085"
ACCENT = "155EEF"
DEEP_BLUE = "1849A9"
CODE_BG = "F5F7FA"
GRID = "D0D5DD"
LATIN_FONT = "Times New Roman"
CJK_FONT = "宋体"


def main() -> int:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    document = Document()

    def color(value: str) -> RGBColor:
        return RGBColor.from_string(value)

    def set_font(style, latin: str, east_asia: str, size: float, *, bold: bool = False) -> None:
        style.font.name = latin
        style.font.size = Pt(size)
        style.font.bold = bold
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:ascii"), latin)
        fonts.set(qn("w:hAnsi"), latin)
        fonts.set(qn("w:eastAsia"), east_asia)
        fonts.set(qn("w:cs"), latin)

    def get_or_create_style(name: str, style_type: WD_STYLE_TYPE):
        try:
            return document.styles[name]
        except KeyError:
            return document.styles.add_style(name, style_type)

    def set_shading(style, fill: str) -> None:
        properties = style.element.get_or_add_pPr()
        shading = properties.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            properties.append(shading)
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), fill)

    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(21)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(22)
    section.right_margin = Mm(22)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    set_font(normal, LATIN_FONT, CJK_FONT, 10.5)
    normal.font.color.rgb = color(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3

    title = document.styles["Title"]
    set_font(title, LATIN_FONT, CJK_FONT, 28, bold=True)
    title.font.color.rgb = color("000000")
    title.paragraph_format.space_before = Pt(110)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    subtitle = document.styles["Subtitle"]
    set_font(subtitle, LATIN_FONT, CJK_FONT, 15)
    subtitle.font.color.rgb = color("000000")
    subtitle.paragraph_format.space_after = Pt(20)
    subtitle.paragraph_format.keep_with_next = True

    heading_specs = {
        "Heading 1": (20, 24, 10),
        "Heading 2": (15.5, 18, 7),
        "Heading 3": (12.5, 14, 5),
        "Heading 4": (11, 11, 4),
    }
    for name, (size, before, after) in heading_specs.items():
        style = document.styles[name]
        set_font(style, LATIN_FONT, CJK_FONT, size, bold=True)
        style.font.color.rgb = color("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True

    body_text = get_or_create_style("Body Text", WD_STYLE_TYPE.PARAGRAPH)
    set_font(body_text, LATIN_FONT, CJK_FONT, 10.5)
    body_text.font.color.rgb = color(INK)
    body_text.paragraph_format.space_after = Pt(6)
    body_text.paragraph_format.line_spacing = 1.3

    list_style = document.styles["List Paragraph"]
    set_font(list_style, LATIN_FONT, CJK_FONT, 10.5)
    list_style.font.color.rgb = color(INK)
    list_style.paragraph_format.space_after = Pt(3)

    code = get_or_create_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    set_font(code, LATIN_FONT, CJK_FONT, 9)
    code.font.color.rgb = color(INK)
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(5)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.05
    code.paragraph_format.keep_together = True
    set_shading(code, CODE_BG)

    inline_code = get_or_create_style("Verbatim Char", WD_STYLE_TYPE.CHARACTER)
    set_font(inline_code, LATIN_FONT, CJK_FONT, 9)
    inline_code.font.color.rgb = color(DEEP_BLUE)

    block_text = get_or_create_style("Block Text", WD_STYLE_TYPE.PARAGRAPH)
    set_font(block_text, LATIN_FONT, CJK_FONT, 10)
    block_text.font.color.rgb = color(MUTED)
    block_text.font.italic = True
    block_text.paragraph_format.left_indent = Inches(0.3)
    block_text.paragraph_format.right_indent = Inches(0.2)
    block_text.paragraph_format.space_before = Pt(5)
    block_text.paragraph_format.space_after = Pt(7)

    caption = document.styles["Caption"]
    set_font(caption, LATIN_FONT, CJK_FONT, 9)
    caption.font.color.rgb = color(MUTED)
    caption.font.italic = False
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)

    table = document.styles["Table Grid"]
    set_font(table, LATIN_FONT, CJK_FONT, 9)
    table.font.color.rgb = color(INK)
    table_properties = table.element.find(qn("w:tblPr"))
    if table_properties is None:
        table_properties = OxmlElement("w:tblPr")
        table.element.append(table_properties)
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), GRID)
        borders.append(border)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", "90"), ("left", "110"), ("bottom", "90"), ("right", "110")):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), value)
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    table_properties.append(margins)

    hyperlink = get_or_create_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    set_font(hyperlink, LATIN_FONT, CJK_FONT, 10.5)
    hyperlink.font.color.rgb = color(ACCENT)
    hyperlink.font.underline = True

    author_style = get_or_create_style("Author", WD_STYLE_TYPE.PARAGRAPH)
    set_font(author_style, LATIN_FONT, CJK_FONT, 10.5)
    author_style.font.color.rgb = color(MUTED)
    author_style.paragraph_format.space_after = Pt(4)

    date_style = get_or_create_style("Date", WD_STYLE_TYPE.PARAGRAPH)
    set_font(date_style, LATIN_FONT, CJK_FONT, 10)
    date_style.font.color.rgb = color(MUTED)
    date_style.paragraph_format.space_after = Pt(6)

    toc_specs = {
        "TOC Heading": (18, 0, 12, True),
        "TOC 1": (11, 0, 4, True),
        "TOC 2": (10.5, 12, 2, False),
        "TOC 3": (9.5, 24, 1, False),
    }
    for name, (size, indent, after, bold) in toc_specs.items():
        style = get_or_create_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_font(style, LATIN_FONT, CJK_FONT, size, bold=bold)
        style.font.color.rgb = color(INK)
        style.paragraph_format.left_indent = Pt(indent)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        if name == "TOC Heading":
            style.paragraph_format.page_break_before = True

    header = section.header
    header.is_linked_to_previous = False
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_p.add_run("NEXENT PRODUCT AND TECHNICAL GUIDE")
    header_run.font.name = LATIN_FONT
    header_run.font.size = Pt(8)
    header_run.font.bold = True
    header_run.font.color.rgb = color("000000")

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run()
    footer_run.font.name = LATIN_FONT
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = color(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        footer_run._r.append(element)

    document.add_paragraph("Nexent Product and Technical Guide", style="Title")
    document.add_paragraph("Deployment Usage Integration and Operations", style="Subtitle")
    metadata = document.add_paragraph("Nexent Project Team")
    metadata.runs[0].font.color.rgb = color(MUTED)
    document.add_section(WD_SECTION.NEW_PAGE)

    sample_heading = document.add_paragraph("Contents", style="Heading 1")
    sample_heading.paragraph_format.page_break_before = False
    document.add_paragraph(
        "This content is replaced by Pandoc. Its presence keeps the reference document valid when opened directly."
    )
    sample_table = document.add_table(rows=2, cols=2, style="Table Grid")
    sample_table.cell(0, 0).text = "Section"
    sample_table.cell(0, 1).text = "Description"
    sample_table.cell(1, 0).text = "Getting started"
    sample_table.cell(1, 1).text = "Install and run Nexent"
    row_properties = sample_table.rows[0]._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row_properties.append(marker)
    for cell in sample_table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), DEEP_BLUE)
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = color("FFFFFF")

    document.core_properties.title = "Nexent Product and Technical Guide"
    document.core_properties.subject = "Customer product and technical documentation"
    document.core_properties.author = "Nexent Project Team"
    document.core_properties.keywords = "Nexent, product guide, technical documentation"
    document.save(OUTPUT_PATH)
    print(f"[build_reference_docx] wrote {OUTPUT_PATH} ({OUTPUT_PATH.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
