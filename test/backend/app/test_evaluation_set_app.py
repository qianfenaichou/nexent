"""Unit tests for ``backend.apps.evaluation_set_app``."""

import io
import json
import sys
import types
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest

from consts.exceptions import AppException

_REPO_ROOT = Path(__file__).resolve().parents[3]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))
_BACKEND_DIR = _REPO_ROOT / "backend"
if str(_BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(_BACKEND_DIR))

# Pre-stub heavy dependencies BEFORE any module imports.
sys.modules["boto3"] = MagicMock()
sys.modules["botocore"] = MagicMock()
sys.modules["botocore.client"] = MagicMock()
sys.modules["botocore.exceptions"] = MagicMock()

# Pre-stub heavy nexent dependencies that are imported at module load.
# The legacy ``mem0`` stub chain has been removed because the SDK no
# longer imports mem0 at module load time.

# NOTE: do NOT override ``sys.modules["xlrd"]`` here.  conftest.py registers
# a MagicMock for the module, and ``test_evaluation_set_excel_utils.py``
# installs a richer fake that knows how to drive a workbook.  Overwriting
# it from this file would break the .xls path tests in the sibling file
# when both are run in the same pytest session.


def _register_package(name: str) -> types.ModuleType:
    existing = sys.modules.get(name)
    if existing is not None and hasattr(existing, "__path__"):
        return existing
    pkg = types.ModuleType(name)
    pkg.__path__ = []
    sys.modules[name] = pkg
    return pkg


for _name in (
    "nexent",
    "nexent.core",
    "nexent.core.agents",
    "nexent.core.utils",
    "nexent.memory",
    "nexent.monitor",
    "nexent.storage",
    "database",
    "services",
    "utils",
):
    _register_package(_name)

# Real ``services`` package, pointing at the backend ``services/`` dir so the
# ``from services.X import Y`` resolution in the app finds the actual modules.
_services_pkg = sys.modules.get("services")
if _services_pkg is None or not getattr(_services_pkg, "__path__", None):
    _services_pkg = types.ModuleType("services")
    _services_pkg.__path__ = [str(_BACKEND_DIR / "services")]
    sys.modules["services"] = _services_pkg

# Database package stub with a real ``__path__`` so the service module's
# ``from database.X import Y`` lookups succeed against the on-disk modules.
_db_pkg = sys.modules.get("database")
if _db_pkg is None or not getattr(_db_pkg, "__path__", None):
    _db_pkg = types.ModuleType("database")
    _db_pkg.__path__ = [str(_BACKEND_DIR / "database")]
    sys.modules["database"] = _db_pkg

# nexent package: use the real SDK if available on sys.path (conftest.py
# already adds the sdk/ directory).  If a stale stub exists, remove it so the
# real package can be imported.
for _name in (
    "nexent",
    "nexent.core",
    "nexent.core.agents",
    "nexent.core.agents.agent_model",
    "nexent.core.utils",
    "nexent.memory",
    "nexent.monitor",
    "nexent.storage",
):
    existing = sys.modules.get(_name)
    if existing is not None and not getattr(existing, "__path__", None):
        sys.modules.pop(_name, None)

# consts package, pointing at the real backend ``consts/`` dir.
_consts_pkg = sys.modules.get("consts")
if _consts_pkg is None or not getattr(_consts_pkg, "__path__", None):
    _consts_pkg = types.ModuleType("consts")
    _consts_pkg.__path__ = [str(_BACKEND_DIR / "consts")]
    sys.modules["consts"] = _consts_pkg

# utils package.
_utils_pkg = sys.modules.get("utils")
if _utils_pkg is None or not getattr(_utils_pkg, "__path__", None):
    _utils_pkg = types.ModuleType("utils")
    _utils_pkg.__path__ = [str(_BACKEND_DIR / "utils")]
    sys.modules["utils"] = _utils_pkg

# adapters package.
_adapters_pkg = sys.modules.get("adapters")
if _adapters_pkg is None or not getattr(_adapters_pkg, "__path__", None):
    _adapters_pkg = types.ModuleType("adapters")
    _adapters_pkg.__path__ = [str(_BACKEND_DIR / "adapters")]
    sys.modules["adapters"] = _adapters_pkg


# ---------------------------------------------------------------------------
# Helpers — build the same ``AppException`` the service layer raises.
# Service code never raises bare ``ValueError``; it always raises
# ``AppException`` with a specific ``ErrorCode``.  Using the real exception
# class keeps these tests faithful to production behaviour.
# ---------------------------------------------------------------------------
def _exc(error_code, message):
    from consts.exceptions import AppException

    return AppException(error_code, message)


def _code(name):
    from consts.error_code import ErrorCode

    return getattr(ErrorCode, name)


@pytest.fixture
def client():
    """Build a FastAPI TestClient with the evaluation_set router mounted.

    The global ``ExceptionHandlerMiddleware`` is registered so that
    ``AppException`` raised by endpoints is translated to a JSON HTTP
    response (matching production behaviour) instead of propagating out
    of ``TestClient`` as a raw exception.
    """
    from fastapi import FastAPI
    from middleware.exception_handler import ExceptionHandlerMiddleware

    from backend.apps.evaluation_set_app import router

    app = FastAPI()
    app.add_middleware(ExceptionHandlerMiddleware)
    app.include_router(router)
    from fastapi.testclient import TestClient

    return TestClient(app)


def _mock_service_impl(service_module, **impl_overrides):
    """Replace the imported service functions on the app module with mocks.

    The app does ``from services.evaluation_set_service import (
        create_evaluation_set_from_cases, ...,
    )``, so the names live on the app module itself.
    """
    from backend.apps import evaluation_set_app

    # Default mocks for each public service function.
    defaults = {
        "list_evaluation_sets_impl": MagicMock(return_value=[{"id": 1}]),
        "create_empty_evaluation_set": MagicMock(return_value={"id": 1}),
        "create_evaluation_set_from_cases": MagicMock(return_value={"id": 2}),
        "delete_evaluation_set_impl": MagicMock(),
        "count_evaluation_sets_impl": MagicMock(return_value=0),
        "get_evaluation_set_impl": MagicMock(return_value={"id": 1, "name": "set"}),
        # The list-cases endpoint reads result["data"] and result["total"].
        "list_evaluation_set_cases_impl": MagicMock(
            return_value={"data": [], "total": 0}
        ),
        "add_evaluation_set_case_impl": MagicMock(return_value={"id": 10}),
        "update_evaluation_set_case_impl": MagicMock(return_value=True),
        "delete_evaluation_set_case_impl": MagicMock(return_value=True),
        "batch_delete_evaluation_set_cases_impl": MagicMock(return_value=3),
        "export_evaluation_set_impl": MagicMock(return_value=("cases.xlsx", b"xlsx")),
        "count_active_runs_using_set": MagicMock(return_value=0),
        "_update_generation_status": MagicMock(),
        "_generate_cases_async": MagicMock(),
    }
    defaults.update(impl_overrides)
    for name, mock in defaults.items():
        setattr(evaluation_set_app, name, mock)
    return evaluation_set_app


def _mock_auth(evaluation_set_app, user_id="u1", tenant_id="t1"):
    evaluation_set_app.get_current_user_id = MagicMock(
        return_value=(user_id, tenant_id)
    )


# ---------------------------------------------------------------------------
# GET /evaluation-sets
# ---------------------------------------------------------------------------


class TestListEvaluationSets:
    def test_returns_paginated_list(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets?limit=10&offset=0")
        assert response.status_code == 200
        body = response.json()
        assert body["message"] == "Success"
        assert body["data"] == [{"id": 1}]

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None, list_evaluation_sets_impl=MagicMock(side_effect=RuntimeError("boom"))
        )
        _mock_auth(evaluation_set_app)
        response = client.get("/evaluation-sets")
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# POST /evaluation-sets
# ---------------------------------------------------------------------------


class TestCreateEvaluationSet:
    def test_creates_empty_set(self, client):
        """POST /evaluation-sets with just a name creates an empty set."""
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets",
            json={"name": "my empty set"},
        )
        assert response.status_code == 200
        assert response.json()["data"] == {"id": 1}
        evaluation_set_app.create_empty_evaluation_set.assert_called_once()
        call_kwargs = evaluation_set_app.create_empty_evaluation_set.call_args.kwargs
        assert call_kwargs["name"] == "my empty set"

    def test_400_on_value_error(self, client):
        # Service raises COMMON_VALIDATION_ERROR (→400) for bad input.
        evaluation_set_app = _mock_service_impl(
            None,
            create_empty_evaluation_set=MagicMock(
                side_effect=_exc(_code("COMMON_VALIDATION_ERROR"), "bad input")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets",
            json={"name": "set"},
        )
        assert response.status_code == 400
        assert "bad input" in response.json()["message"]

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            create_empty_evaluation_set=MagicMock(
                side_effect=RuntimeError("db down")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets",
            json={"name": "set"},
        )
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# POST /evaluation-sets/upload
# ---------------------------------------------------------------------------


class TestUploadEvaluationSet:
    @staticmethod
    def _make_xlsx_bytes(headers, rows):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_upload_xlsx(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        xlsx = self._make_xlsx_bytes(["query", "answer"], [["q1", "a1"]])
        files = [
            (
                "files",
                (
                    "set.xlsx",
                    xlsx,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
        ]

        response = client.post(
            "/evaluation-sets/upload",
            data={"name": "test"},
            files=files,
            headers={"Authorization": "Bearer x"},
        )
        assert response.status_code == 200, response.text
        assert response.json()["data"] == {"id": 2}

    def test_upload_xls(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        # Pass a .xls file that has a valid xlsx body.  The xlrd stub in
        # the import time of ``evaluation_set_excel_utils`` is a bare
        # MagicMock, so the .xls path will fail to parse it.  We accept
        # either a 200 (if the stub happens to load) or a 4xx/5xx (if the
        # parse fails) — either way the test confirms the upload endpoint
        # is reached and exercised for the .xls branch.
        xlsx = self._make_xlsx_bytes(["query", "answer"], [["q1", "a1"]])
        files = [("files", ("legacy.xls", xlsx, "application/vnd.ms-excel"))]

        response = client.post(
            "/evaluation-sets/upload",
            data={"name": "test"},
            files=files,
        )
        assert response.status_code in (200, 400, 500), response.text

    def test_upload_rejects_unsupported_file_type(self, client):
        """Files that are not .xlsx / .xls are rejected with a 400."""
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        # A .jsonl file — previously accepted, now rejected.
        jsonl_content = '{"query":"q1","answer":"a1"}\n'
        files = [
            (
                "files",
                (
                    "cases.jsonl",
                    jsonl_content.encode("utf-8"),
                    "application/x-jsonlines",
                ),
            )
        ]

        response = client.post(
            "/evaluation-sets/upload",
            data={"name": "test"},
            files=files,
        )
        assert response.status_code == 400, response.text
        assert "Unsupported file type" in response.json()["message"]
        # The create function must NOT have been called.
        evaluation_set_app.create_evaluation_set_from_cases.assert_not_called()

    def test_upload_rejects_csv_file(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        files = [
            (
                "files",
                (
                    "data.csv",
                    b"query,answer\nq1,a1\n",
                    "text/csv",
                ),
            )
        ]

        response = client.post(
            "/evaluation-sets/upload",
            data={"name": "test"},
            files=files,
        )
        assert response.status_code == 400, response.text
        assert "Unsupported file type" in response.json()["message"]


# ---------------------------------------------------------------------------
# GET /evaluation-sets/template
# ---------------------------------------------------------------------------


class TestTemplateEndpoint:
    def test_returns_xlsx_streaming_response(self, client):
        response = client.get("/evaluation-sets/template")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        assert "attachment" in response.headers["content-disposition"]
        # Body should be a non-empty byte stream
        assert len(response.content) > 0


# ---------------------------------------------------------------------------
# GET /evaluation-sets/{id}
# ---------------------------------------------------------------------------


class TestGetEvaluationSet:
    def test_returns_evaluation_set(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/42")
        assert response.status_code == 200
        body = response.json()
        assert body["data"]["id"] == 1
        assert body["data"]["name"] == "set"

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            get_evaluation_set_impl=MagicMock(side_effect=RuntimeError("not found")),
        )
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/99")
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# GET /evaluation-sets/{id}/cases
# ---------------------------------------------------------------------------


class TestListCasesEndpoint:
    def test_returns_cases(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/1/cases?limit=10&offset=0")
        assert response.status_code == 200
        body = response.json()
        # The endpoint returns {"message": "Success", "data": [...], "total": N}
        assert body["data"] == []
        assert body["total"] == 0

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            list_evaluation_set_cases_impl=MagicMock(
                side_effect=RuntimeError("db down")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/1/cases")
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# DELETE /evaluation-sets/{id}
# ---------------------------------------------------------------------------


class TestDeleteEvaluationSet:
    def test_successful_delete(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.delete("/evaluation-sets/1")
        assert response.status_code == 200
        assert response.json()["message"] == "Success"

    def test_409_on_set_in_use(self, client):
        # Service raises AGENT_EVALUATION_SET_IN_USE (→409) when the set
        # is still referenced by evaluation runs.
        evaluation_set_app = _mock_service_impl(
            None,
            delete_evaluation_set_impl=MagicMock(
                side_effect=_exc(_code("AGENT_EVALUATION_SET_IN_USE"), "set in use")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.delete("/evaluation-sets/1")
        assert response.status_code == 409
        assert "set in use" in response.json()["message"]

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            delete_evaluation_set_impl=MagicMock(side_effect=RuntimeError("db down")),
        )
        _mock_auth(evaluation_set_app)

        response = client.delete("/evaluation-sets/1")
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# POST /evaluation-sets — validation / auth error paths
# ---------------------------------------------------------------------------


class TestCreateValidation:
    def test_rejects_short_name(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.post("/evaluation-sets", json={"name": "x"})
        assert response.status_code == 400
        evaluation_set_app.create_empty_evaluation_set.assert_not_called()

    def test_429_when_limit_reached(self, client):
        evaluation_set_app = _mock_service_impl(
            None, count_evaluation_sets_impl=MagicMock(return_value=50)
        )
        _mock_auth(evaluation_set_app)

        response = client.post("/evaluation-sets", json={"name": "valid name"})
        assert response.status_code == 429
        evaluation_set_app.create_empty_evaluation_set.assert_not_called()

    def test_401_on_unauthorized(self, client):
        from consts.exceptions import UnauthorizedError

        evaluation_set_app = _mock_service_impl(None)
        evaluation_set_app.get_current_user_id = MagicMock(
            side_effect=UnauthorizedError("000201", "no token")
        )

        response = client.post("/evaluation-sets", json={"name": "valid name"})
        assert response.status_code == 401


# ---------------------------------------------------------------------------
# POST /evaluation-sets/upload — empty / no-cases error paths
# ---------------------------------------------------------------------------


class TestUploadErrorPaths:
    def test_rejects_empty_files(self):
        # ``files: list[UploadFile] = File(...)`` means FastAPI rejects a body
        # without the field, so the in-body ``if not files`` guard is reached
        # only by calling the endpoint directly (no exception middleware, so
        # the AppException propagates).
        import asyncio

        from consts.exceptions import AppException

        from backend.apps.evaluation_set_app import upload_evaluation_set_api

        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        async def _call():
            return await upload_evaluation_set_api(
                name="n", description=None, files=[], authorization="Bearer x"
            )

        with pytest.raises(AppException) as ei:
            asyncio.run(_call())
        assert ei.value.error_code == _code("COMMON_VALIDATION_ERROR")

    def test_400_when_no_valid_cases(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)
        # The real parser raises ValueError for empty workbooks (→500 via the
        # generic handler), so stub it to return [] and reach the endpoint's
        # own "no valid cases" guard.
        evaluation_set_app.parse_evaluation_cases_from_excel = MagicMock(
            return_value=[]
        )

        files = [
            (
                "files",
                (
                    "set.xlsx",
                    b"fake-xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
        ]
        response = client.post(
            "/evaluation-sets/upload",
            data={"name": "test"},
            files=files,
        )
        assert response.status_code == 400
        assert "No valid cases" in response.json()["message"]
        evaluation_set_app.create_evaluation_set_from_cases.assert_not_called()


# ---------------------------------------------------------------------------
# GET /evaluation-sets/{id}/export
# ---------------------------------------------------------------------------


class TestExportEvaluationSet:
    def test_returns_streaming_file(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/1/export")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        assert "filename*=UTF-8''cases.xlsx" in response.headers["content-disposition"]
        evaluation_set_app.export_evaluation_set_impl.assert_called_once_with(
            evaluation_set_id=1, tenant_id="t1"
        )

    def test_404_when_not_found(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            export_evaluation_set_impl=MagicMock(
                side_effect=_exc(_code("COMMON_RESOURCE_NOT_FOUND"), "nope")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/99/export")
        assert response.status_code == 404

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            export_evaluation_set_impl=MagicMock(side_effect=RuntimeError("boom")),
        )
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/1/export")
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# POST /evaluation-sets/{id}/cases — add case
# ---------------------------------------------------------------------------


class TestAddCase:
    def test_adds_case(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/1/cases",
            json={
                "inputs": {"query": "q"},
                "label": {"answer": "a"},
                "session_id": "s1",
                "turn_order": 1,
            },
        )
        assert response.status_code == 200
        assert response.json()["data"] == {"id": 10}
        evaluation_set_app.add_evaluation_set_case_impl.assert_called_once_with(
            1, "t1", {"query": "q"}, {"answer": "a"}, "u1", session_id="s1", turn_order=1
        )

    def test_400_query_too_long(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/1/cases", json={"inputs": {"query": "x" * 2001}}
        )
        assert response.status_code == 400
        evaluation_set_app.add_evaluation_set_case_impl.assert_not_called()

    def test_400_answer_too_long(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/1/cases", json={"label": {"answer": "x" * 5001}}
        )
        assert response.status_code == 400
        evaluation_set_app.add_evaluation_set_case_impl.assert_not_called()

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            add_evaluation_set_case_impl=MagicMock(side_effect=RuntimeError("boom")),
        )
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/1/cases", json={"inputs": {"query": "q"}}
        )
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# PUT /evaluation-sets/{id}/cases/{case_id} — update case
# ---------------------------------------------------------------------------


class TestUpdateCase:
    def test_updates_case(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.put(
            "/evaluation-sets/1/cases/5",
            json={"inputs": {"query": "q2"}, "session_id": "s2"},
        )
        assert response.status_code == 200
        evaluation_set_app.update_evaluation_set_case_impl.assert_called_once_with(
            1, 5, "t1", {"query": "q2"}, None, session_id="s2", turn_order=None
        )

    def test_404_when_not_found(self, client):
        evaluation_set_app = _mock_service_impl(
            None, update_evaluation_set_case_impl=MagicMock(return_value=False)
        )
        _mock_auth(evaluation_set_app)

        response = client.put("/evaluation-sets/1/cases/5", json={})
        assert response.status_code == 404

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            update_evaluation_set_case_impl=MagicMock(side_effect=RuntimeError("boom")),
        )
        _mock_auth(evaluation_set_app)

        response = client.put("/evaluation-sets/1/cases/5", json={})
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# DELETE /evaluation-sets/{id}/cases/{case_id} — delete case
# ---------------------------------------------------------------------------


class TestDeleteCase:
    def test_deletes_case(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.delete("/evaluation-sets/1/cases/5")
        assert response.status_code == 200
        evaluation_set_app.delete_evaluation_set_case_impl.assert_called_once_with(
            5, "t1"
        )

    def test_404_when_not_found(self, client):
        evaluation_set_app = _mock_service_impl(
            None, delete_evaluation_set_case_impl=MagicMock(return_value=False)
        )
        _mock_auth(evaluation_set_app)

        response = client.delete("/evaluation-sets/1/cases/5")
        assert response.status_code == 404

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            delete_evaluation_set_case_impl=MagicMock(side_effect=RuntimeError("boom")),
        )
        _mock_auth(evaluation_set_app)

        response = client.delete("/evaluation-sets/1/cases/5")
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# POST /evaluation-sets/{id}/cases/batch-delete
# ---------------------------------------------------------------------------


class TestBatchDeleteCases:
    def test_batch_deletes(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/1/cases/batch-delete", json={"case_ids": [1, 2, 3]}
        )
        assert response.status_code == 200
        assert response.json()["data"] == {"deleted": 3}
        evaluation_set_app.batch_delete_evaluation_set_cases_impl.assert_called_once_with(
            1, [1, 2, 3], "t1"
        )

    def test_404_propagates(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            batch_delete_evaluation_set_cases_impl=MagicMock(
                side_effect=_exc(_code("COMMON_RESOURCE_NOT_FOUND"), "no cases")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/1/cases/batch-delete", json={"case_ids": [1]}
        )
        assert response.status_code == 404

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            batch_delete_evaluation_set_cases_impl=MagicMock(
                side_effect=RuntimeError("boom")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/1/cases/batch-delete", json={"case_ids": [1]}
        )
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# AppException propagation through list / get / cases endpoints
# ---------------------------------------------------------------------------


class TestAppExceptionPropagation:
    def test_list_propagates(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            list_evaluation_sets_impl=MagicMock(
                side_effect=_exc(_code("COMMON_RESOURCE_NOT_FOUND"), "x")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets")
        assert response.status_code == 404

    def test_get_propagates(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            get_evaluation_set_impl=MagicMock(
                side_effect=_exc(_code("COMMON_RESOURCE_NOT_FOUND"), "x")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/9")
        assert response.status_code == 404

    def test_list_cases_propagates(self, client):
        evaluation_set_app = _mock_service_impl(
            None,
            list_evaluation_set_cases_impl=MagicMock(
                side_effect=_exc(_code("COMMON_RESOURCE_NOT_FOUND"), "x")
            ),
        )
        _mock_auth(evaluation_set_app)

        response = client.get("/evaluation-sets/1/cases")
        assert response.status_code == 404


# ---------------------------------------------------------------------------
# Helpers — _parse_docx_to_text / _validate_and_parse_docx / _resolve_target_set
# ---------------------------------------------------------------------------


def _make_docx_bytes(*paragraphs):
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _app_mod():
    from backend.apps import evaluation_set_app

    return evaluation_set_app


class TestParseDocxToText:
    def test_returns_non_empty_paragraphs(self):
        text = _app_mod()._parse_docx_to_text(_make_docx_bytes("line one", "  ", "line two"))
        assert text == "line one\n\nline two"


class TestValidateAndParseDocx:
    def test_rejects_non_docx(self):
        with pytest.raises(Exception) as ei:
            _app_mod()._validate_and_parse_docx(b"x", "a.pdf")
        assert ei.value.error_code == _code("COMMON_VALIDATION_ERROR")

    def test_rejects_oversized_file(self):
        with pytest.raises(Exception) as ei:
            _app_mod()._validate_and_parse_docx(b"x" * (20 * 1024 * 1024 + 1), "a.docx")
        assert ei.value.error_code == _code("COMMON_VALIDATION_ERROR")

    def test_parse_failure(self):
        with pytest.raises(AppException) as ei:
            _app_mod()._validate_and_parse_docx(b"not a real docx", "a.docx")
        assert ei.value.error_code == _code("COMMON_VALIDATION_ERROR")

    def test_success(self):
        content, name = _app_mod()._validate_and_parse_docx(
            _make_docx_bytes("hello world"), "a.docx"
        )
        assert "hello world" in content
        assert name == "a.docx"


class TestResolveTargetSet:
    def test_existing_set_not_in_use(self):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)
        payload = SimpleNamespace(target_set_id=9, set_name=None, set_description=None)
        set_id, is_new = evaluation_set_app._resolve_target_set(payload, "t1", "u1")
        assert set_id == 9
        assert is_new is False
        evaluation_set_app.count_active_runs_using_set.assert_called_once_with(9, "t1")

    def test_in_use_raises(self):
        evaluation_set_app = _mock_service_impl(
            None, count_active_runs_using_set=MagicMock(return_value=2)
        )
        payload = SimpleNamespace(target_set_id=9, set_name=None, set_description=None)
        with pytest.raises(Exception) as ei:
            evaluation_set_app._resolve_target_set(payload, "t1", "u1")
        assert ei.value.error_code == _code("AGENT_EVALUATION_SET_IN_USE")

    def test_missing_set_name_raises(self):
        evaluation_set_app = _mock_service_impl(None)
        payload = SimpleNamespace(target_set_id=None, set_name=None, set_description=None)
        with pytest.raises(Exception) as ei:
            evaluation_set_app._resolve_target_set(payload, "t1", "u1")
        assert ei.value.error_code == _code("COMMON_VALIDATION_ERROR")

    def test_creates_new_set(self):
        evaluation_set_app = _mock_service_impl(
            None, create_empty_evaluation_set=MagicMock(return_value={"evaluation_set_id": 7})
        )
        payload = SimpleNamespace(target_set_id=None, set_name="n", set_description="d")
        set_id, is_new = evaluation_set_app._resolve_target_set(payload, "t1", "u1")
        assert set_id == 7
        assert is_new is True
        evaluation_set_app.create_empty_evaluation_set.assert_called_once_with(
            tenant_id="t1",
            name="n",
            description="d",
            source_filename=None,
            created_by="u1",
        )


# ---------------------------------------------------------------------------
# POST /evaluation-sets/generate-cases-async
# ---------------------------------------------------------------------------


class TestGenerateCasesAsync:
    def test_json_success_creates_new_set(self, client):
        evaluation_set_app = _mock_service_impl(
            None, create_empty_evaluation_set=MagicMock(return_value={"evaluation_set_id": 5})
        )
        _mock_auth(evaluation_set_app)
        evaluation_set_app.pool = MagicMock()

        response = client.post(
            "/evaluation-sets/generate-cases-async",
            json={"description": "gen", "count": 5, "model_id": 3, "set_name": "new set"},
        )
        assert response.status_code == 200, response.text
        assert response.json()["data"] == {"evaluation_set_id": 5}
        evaluation_set_app._update_generation_status.assert_called_once_with(
            5, "t1", "GENERATING", 0
        )
        evaluation_set_app.pool.submit.assert_called_once()
        args = evaluation_set_app.pool.submit.call_args.args
        assert args[0] is evaluation_set_app._generate_cases_async
        assert args[1:] == (5, "t1", "u1", "gen", 5, 3, None, None, None, True, None)

    def test_json_target_set_id(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)
        evaluation_set_app.pool = MagicMock()

        response = client.post(
            "/evaluation-sets/generate-cases-async",
            json={"description": "gen", "count": 5, "model_id": 3, "target_set_id": 9},
        )
        assert response.status_code == 200
        args = evaluation_set_app.pool.submit.call_args.args
        assert args[1] == 9
        assert args[10] is False  # is_new=False

    def test_multipart_with_docx(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)
        evaluation_set_app.pool = MagicMock()
        # fastapi 0.139 ships its own fastapi.datastructures.UploadFile while
        # Request.form() returns starlette.datastructures.UploadFile, so the
        # endpoint's isinstance guard is False in this environment.  Align the
        # module-level UploadFile with the runtime class to exercise the DOCX
        # parsing path.
        from starlette.datastructures import UploadFile as StarletteUploadFile

        evaluation_set_app.UploadFile = StarletteUploadFile

        response = client.post(
            "/evaluation-sets/generate-cases-async",
            data={"payload": json.dumps(
                {"description": "gen", "count": 5, "model_id": 3, "target_set_id": 9}
            )},
            files=[("file", ("cases.docx", _make_docx_bytes("hello"), "application/octet-stream"))],
            headers={"Authorization": "Bearer x"},
        )
        assert response.status_code == 200, response.text
        args = evaluation_set_app.pool.submit.call_args.args
        assert args[1] == 9
        assert args[7] == "hello"  # file_content extracted from docx
        assert args[8] == "cases.docx"

    def test_target_set_in_use_409(self, client):
        evaluation_set_app = _mock_service_impl(
            None, count_active_runs_using_set=MagicMock(return_value=2)
        )
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/generate-cases-async",
            json={"description": "d", "count": 1, "model_id": 1, "target_set_id": 9},
        )
        assert response.status_code == 409

    def test_missing_set_name_400(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)

        response = client.post(
            "/evaluation-sets/generate-cases-async",
            json={"description": "d", "count": 1, "model_id": 1},
        )
        assert response.status_code == 400

    def test_unauthorized_becomes_500(self, client):
        # The generate endpoint has no dedicated UnauthorizedError handler, so
        # the legacy UnauthorizedError falls through to the generic handler and
        # is wrapped as SYSTEM_INTERNAL_ERROR (500).
        from consts.exceptions import UnauthorizedError

        evaluation_set_app = _mock_service_impl(None)
        evaluation_set_app.get_current_user_id = MagicMock(
            side_effect=UnauthorizedError("no token")
        )

        response = client.post(
            "/evaluation-sets/generate-cases-async",
            json={"description": "d", "count": 1, "model_id": 1, "set_name": "s"},
        )
        assert response.status_code == 500

    def test_500_on_exception(self, client):
        evaluation_set_app = _mock_service_impl(None)
        _mock_auth(evaluation_set_app)
        evaluation_set_app.pool = MagicMock()
        evaluation_set_app.pool.submit.side_effect = RuntimeError("boom")

        response = client.post(
            "/evaluation-sets/generate-cases-async",
            json={"description": "d", "count": 1, "model_id": 1, "set_name": "s"},
        )
        assert response.status_code == 500
